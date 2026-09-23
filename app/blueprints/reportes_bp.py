import os
import json
import time
import io
import datetime
from flask import Blueprint, render_template, request, jsonify, send_file, session, redirect, url_for, flash
from ..db import get_db_connection
from .auth import admin_required, solo_admin
from ..reportes import CATALOGO

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

bp = Blueprint('reportes', __name__)

REPORTES_INFO = {
    'consulta_cuentas': {
        'nombre': 'Consulta de Cuentas',
        'categoria': 'Contabilidad',
        'icono': 'bi-journal-text',
        'descripcion': 'REG_CTAS agrupado por tercero con débito, crédito y saldo neto.',
        'template': 'reporte_consulta_cuentas.html',
    },
    'consulta_tercero': {
        'nombre': 'Consulta por Tercero',
        'categoria': 'Contabilidad',
        'icono': 'bi-person-lines-fill',
        'descripcion': 'Movimientos detallados de cuentas por tercero.',
        'template': 'reporte_consulta_tercero.html',
    },
    'reg_ctas': {
        'nombre': 'Registro de Cuentas',
        'categoria': 'Contabilidad',
        'icono': 'bi-table',
        'descripcion': 'Listado completo de asientos y registros contables.',
        'template': 'reporte_reg_ctas.html',
    },
    'matriz_terceros_cuentas': {
        'nombre': 'Matriz Terceros vs Cuentas',
        'categoria': 'Contabilidad',
        'icono': 'bi-grid-3x3',
        'descripcion': 'Matriz cruzada de terceros y cuentas contables.',
        'template': 'reporte_matriz_terceros_cuentas.html',
    },
    'ventas_clientes': {
        'nombre': 'Ventas por Cliente',
        'categoria': 'Ventas',
        'icono': 'bi-cart-check',
        'descripcion': 'Consolidado de ventas facturadas por cliente.',
        'template': 'reporte_ventas_clientes.html',
    },
    'inventario': {
        'nombre': 'Inventario y Existencias',
        'categoria': 'Inventario',
        'icono': 'bi-boxes',
        'descripcion': 'Existencias, costos y saldos de productos.',
        'template': 'reporte_inventario.html',
    },
    'auditoria_facturas': {
        'nombre': 'Auditoría de Documentos de Interfase',
        'categoria': 'Contabilidad',
        'icono': 'bi-shield-check',
        'descripcion': 'Auditoría y validación contable de documentos y facturas generadas mediante interfases con otros sistemas.',
        'template': 'reporte_auditoria_facturas.html',
    },
}

def _obtener_agentes_activos(conn, usuario_id, rol):
    try:
        rows = conn.execute("""
            SELECT cliente_id, nombre, alias, ultimo_ping, version, db_version, ruta_bd
            FROM admin_agent_sesiones
            WHERE ultimo_ping > NOW() - INTERVAL '30 days'
            ORDER BY ultimo_ping DESC, id DESC
        """).fetchall()
    except Exception:
        rows = conn.execute("""
            SELECT cliente_id, nombre, alias, ultimo_ping, version
            FROM admin_agent_sesiones
            WHERE ultimo_ping > NOW() - INTERVAL '30 days'
            ORDER BY ultimo_ping DESC, id DESC
        """).fetchall()
    
    permitidos = None
    if rol != 'Administrador':
        perm_rows = conn.execute("""
            SELECT cliente_id FROM admin_agent_permisos WHERE usuario_id = %s
        """, (usuario_id,)).fetchall()
        permitidos = {r['cliente_id'].lower() for r in perm_rows}

    now = datetime.datetime.now(datetime.timezone.utc)
    maquinas = {}
    for r in rows:
        cid_raw = r['cliente_id'] or ''
        cid_clean = cid_raw.strip().lower()
        base_key = cid_clean.replace('_daemon', '')
        if not base_key:
            continue
        if permitidos is not None and base_key not in permitidos and cid_clean not in permitidos:
            continue
            
        up = r['ultimo_ping']
        if up and up.tzinfo is None:
            up = up.replace(tzinfo=datetime.timezone.utc)
            
        diff = (now - up).total_seconds() if up else 9999999
        online = diff < 300
        
        alias = (r['alias'] or '').strip()
        nombre = (r['nombre'] or '').strip()
        ver = (r.get('version') or '').strip()
        db_ver = (r.get('db_version') or '').strip()
        ruta_bd = (r.get('ruta_bd') or '').strip()
        if ver and not ver.lower().startswith('v'):
            ver = f'v{ver}'
        if nombre.endswith('(Daemon)'):
            nombre = nombre[:-8].strip()
            
        display_name = alias or nombre or base_key
        
        if base_key not in maquinas:
            maquinas[base_key] = {
                'id': base_key,
                'nombre': display_name,
                'alias': alias,
                'version': ver,
                'db_version': db_ver,
                'ruta_bd': ruta_bd,
                'online': online,
                'max_up': up,
                'min_diff': diff,
            }
        else:
            if online:
                maquinas[base_key]['online'] = True
            if alias and not maquinas[base_key]['alias']:
                maquinas[base_key]['alias'] = alias
                maquinas[base_key]['nombre'] = alias
            if ver and not maquinas[base_key]['version']:
                maquinas[base_key]['version'] = ver
            if db_ver and not maquinas[base_key].get('db_version'):
                maquinas[base_key]['db_version'] = db_ver
            if ruta_bd and not maquinas[base_key].get('ruta_bd'):
                maquinas[base_key]['ruta_bd'] = ruta_bd
            if up and (maquinas[base_key]['max_up'] is None or up > maquinas[base_key]['max_up']):
                maquinas[base_key]['max_up'] = up
                maquinas[base_key]['min_diff'] = diff

    agentes = []
    for base_key, m in maquinas.items():
        diff = m['min_diff']
        mins = int(diff / 60)
        if mins < 60:
            ultimo = f'hace {mins}m'
        elif mins < 1440:
            ultimo = f'hace {mins // 60}h'
        else:
            ultimo = f'hace {mins // 1440}d'
            
        agentes.append({
            'id': m['id'],
            'nombre': m['nombre'],
            'alias': m['alias'],
            'version': m['version'] or 'v1.2.3',
            'db_version': m.get('db_version') or '',
            'ruta_bd': m.get('ruta_bd') or '',
            'online': m['online'],
            'ultimo': ultimo,
            '_diff': diff
        })
        
    agentes.sort(key=lambda a: (not a['online'], a['_diff']))
    for a in agentes:
        a.pop('_diff', None)
    return agentes

def _ejecutar_consulta_remota(conn, cliente_id, tipo, parametros, timeout=90):
    row = conn.execute("""
        INSERT INTO admin_agent_consultas (sesion_id, tipo, parametros, estado)
        SELECT id, %s, %s::jsonb, 'pendiente'
        FROM admin_agent_sesiones
        WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon')
           OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '')
        ORDER BY ultimo_ping DESC, id DESC
        LIMIT 1
        RETURNING id
    """, (tipo, json.dumps(parametros), cliente_id, cliente_id, cliente_id)).fetchone()
    
    if not row:
        row = conn.execute("""
            INSERT INTO admin_agent_consultas (sesion_id, tipo, parametros, estado)
            SELECT id, %s, %s::jsonb, 'pendiente'
            FROM admin_agent_sesiones
            WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon')
               OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '')
            ORDER BY id DESC
            LIMIT 1
            RETURNING id
        """, (tipo, json.dumps(parametros), cliente_id, cliente_id, cliente_id)).fetchone()
    
    if not row:
        raise RuntimeError(f"El agente '{cliente_id}' no está registrado.")
        
    conn.commit()
    consulta_id = row['id']
    
    t0 = time.time()
    while time.time() - t0 < timeout:
        time.sleep(0.5)
        c = conn.execute("""
            SELECT estado, respuesta FROM admin_agent_consultas WHERE id=%s
        """, (consulta_id,)).fetchone()
        if c:
            if c['estado'] == 'lista':
                return c['respuesta']
            if c['estado'] == 'error':
                err = c.get('respuesta')
                if isinstance(err, dict):
                    raise RuntimeError(err.get('error', 'Error en el agente'))
                raise RuntimeError(str(err or 'Error en el agente'))
    raise TimeoutError(f"El agente '{cliente_id}' tardó más de {timeout}s en responder.")

def _resolver_terceros_lote(conn, agente, cod_ters):
    """Resuelve en lote una lista de códigos de terceros (COD_TER) usando la caché en PostgreSQL.
    Si la caché no existe o hay códigos faltantes, sincroniza la tabla TERCEROS desde el agente una sola vez."""
    if not cod_ters:
        return {}
    clean_ids = list({str(c).strip() for c in cod_ters if str(c).strip() and str(c).strip() != '0'})
    if not clean_ids:
        return {}

    agente_clean = agente.strip().lower().replace('_daemon', '')
    try:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS admin_terceros_cache (
                id SERIAL PRIMARY KEY,
                cliente_id VARCHAR(100) NOT NULL,
                cod_ter VARCHAR(50) NOT NULL,
                nombre VARCHAR(250),
                nit VARCHAR(50),
                updated_at TIMESTAMPTZ DEFAULT NOW(),
                UNIQUE(cliente_id, cod_ter)
            )
        """)
        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_terceros_cache_lookup ON admin_terceros_cache (LOWER(cliente_id), cod_ter)
        """)
        conn.commit()

        # 1. Buscar en cache local
        rows = conn.execute("""
            SELECT cod_ter, nombre, nit FROM admin_terceros_cache
            WHERE LOWER(cliente_id) = LOWER(%s) AND cod_ter = ANY(%s)
        """, (agente_clean, clean_ids)).fetchall()

        mapa = {r['cod_ter']: {'nombre': r['nombre'] or '', 'nit': r['nit'] or ''} for r in rows}
        faltantes = [cid for cid in clean_ids if cid not in mapa]

        # 2. Si hay faltantes o la caché está vacía, consultar TERCEROS al agente y poblar la caché
        cnt_row = conn.execute("SELECT COUNT(*) AS c FROM admin_terceros_cache WHERE LOWER(cliente_id)=LOWER(%s)", (agente_clean,)).fetchone()
        total_en_cache = cnt_row['c'] if cnt_row else 0

        if total_en_cache == 0 or len(faltantes) > 0:
            try:
                datos_ter = _ejecutar_consulta_remota(conn, agente, 'multi_tabla', {
                    'tablas': [{'tabla': 'TERCEROS', 'campos': ['COD_TER', 'NOMBRE', 'NIT', 'IDENTIFICA'], 'filtros': {}}]
                }, timeout=90)
                raw_ter = datos_ter.get('TERCEROS', []) if isinstance(datos_ter, dict) else []
                if isinstance(raw_ter, list) and len(raw_ter) > 0:
                    insert_tuples = []
                    for r in raw_ter:
                        c = str(r.get('COD_TER', '') or '').strip()
                        n = str(r.get('NOMBRE', '') or '').strip()
                        nit = str(r.get('IDENTIFICA') or r.get('NIT') or '').strip()
                        if c:
                            insert_tuples.append((agente_clean, c, n, nit))
                            if c in clean_ids:
                                mapa[c] = {'nombre': n, 'nit': nit}

                    if insert_tuples:
                        try:
                            import psycopg2.extras
                            raw_cur = conn._conn.cursor() if hasattr(conn, '_conn') else conn.cursor()
                            psycopg2.extras.execute_values(
                                raw_cur,
                                """
                                INSERT INTO admin_terceros_cache (cliente_id, cod_ter, nombre, nit, updated_at)
                                VALUES %s
                                ON CONFLICT (cliente_id, cod_ter)
                                DO UPDATE SET nombre=EXCLUDED.nombre, nit=EXCLUDED.nit, updated_at=NOW()
                                """,
                                insert_tuples,
                                page_size=2000
                            )
                            conn.commit()
                        except Exception as e_ins:
                            print(f"Fallback insert terceros cache: {e_ins}")
            except Exception as ex:
                print(f"Advertencia sincronizando terceros en cache desde agente: {ex}")

        return mapa
    except Exception as e:
        print(f"Error en _resolver_terceros_lote: {e}")
        return {}

def _generar_excel_buffer(reporte_id, rows, filtros, modulo=None, fuente='remoto', cliente_id=''):
    if not HAS_OPENPYXL:
        raise RuntimeError('openpyxl no está instalado')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = reporte_id[:31]

    azul_oscuro = '003F7F'
    gris_claro  = 'F0F4FF'
    hdr_font    = Font(bold=True, color='FFFFFF')
    hdr_fill    = PatternFill('solid', fgColor=azul_oscuro)
    tit_font    = Font(bold=True, size=12, color='FFFFFF')
    tit_fill    = PatternFill('solid', fgColor=azul_oscuro)
    meta_font   = Font(italic=True, color='555555', size=10)
    pie_font    = Font(italic=True, color='888888', size=9)
    num_fmt     = '#,##0.00'

    fila = 1
    nombre_reporte = getattr(modulo, 'NOMBRE', reporte_id) if modulo else reporte_id
    c = ws.cell(fila, 1, nombre_reporte)
    c.font = tit_font
    c.fill = tit_fill
    fila += 1

    fuente_txt = f"Fuente: {fuente.upper()}"
    if cliente_id:
        fuente_txt += f"  —  {cliente_id}"
    c = ws.cell(fila, 1, fuente_txt)
    c.font = meta_font
    fila += 1

    filtros_ui = getattr(modulo, 'FILTROS_UI', []) if modulo else []
    for f_ui in filtros_ui:
        fid = f_ui['id']
        val = filtros.get(fid, '')
        if isinstance(val, list):
            val = ', '.join(str(v) for v in val if v)
        if val:
            c = ws.cell(fila, 1, f"{f_ui['label']}: {val}")
            c.font = meta_font
            fila += 1

    fila += 1

    if not rows:
        ws.cell(fila, 1, 'Sin resultados').font = meta_font
    else:
        excel_cols = getattr(modulo, 'EXCEL_COLS', None) if modulo else None
        if excel_cols:
            keys   = [k  for k, _ in excel_cols]
            labels = [lb for _, lb in excel_cols]
        else:
            keys   = list(rows[0].keys())
            labels = [k.replace('_', ' ').title() for k in keys]

        for ci, label in enumerate(labels, 1):
            c = ws.cell(fila, ci, label)
            c.font = hdr_font
            c.fill = hdr_fill
        fila += 1

        sample = rows[0]
        num_keys = {k for k in keys if isinstance(sample.get(k), (int, float))}

        for ri, r in enumerate(rows):
            fill_par = PatternFill('solid', fgColor=gris_claro) if ri % 2 == 1 else None
            for ci, k in enumerate(keys, 1):
                val = r.get(k, '')
                c = ws.cell(fila, ci, val)
                if k in num_keys:
                    c.number_format = num_fmt
                if fill_par:
                    c.fill = fill_par
            fila += 1

    fila += 1
    ws.cell(fila, 1, f'Generado: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}').font = pie_font

    for col in ws.columns:
        letra = col[0].column_letter
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        ws.column_dimensions[letra].width = min(max_len + 2, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

@bp.route('/admin/reportes')
@admin_required
def reportes_menu():
    tipo = request.args.get('tipo') or request.args.get('id') or request.args.get('reporte')
    if tipo:
        return reporte_detalle(tipo)
    conn = get_db_connection()
    try:
        usuario_id = session.get('usuario_id')
        rol = session.get('rol', '')
        agentes = _obtener_agentes_activos(conn, usuario_id, rol)
        
        categorias = {}
        for rid, info in REPORTES_INFO.items():
            cat = info['categoria']
            if cat not in categorias:
                categorias[cat] = []
            categorias[cat].append({
                'id': rid,
                'nombre': info['nombre'],
                'icono': info.get('icono', 'bi-file-earmark-text'),
                'descripcion': info.get('descripcion', ''),
            })
            
        agente_param = request.args.get('agente', '').strip()
        if not agente_param and agentes:
            online_agents = [a for a in agentes if a['online']]
            agente_param = online_agents[0]['id'] if online_agents else agentes[0]['id']

        return render_template('sar_reportes_menu.html',
                               categorias=categorias,
                               agentes=agentes,
                               agente_actual=agente_param,
                               nombre=session.get('nombre', ''))
    finally:
        conn.close()

@bp.route('/admin/reportes/<reporte_id>')
@admin_required
def reporte_detalle(reporte_id):
    if reporte_id not in REPORTES_INFO and reporte_id not in CATALOGO:
        flash(f"Reporte '{reporte_id}' no encontrado", "warning")
        return redirect(url_for('reportes.reportes_menu'))
        
    info = REPORTES_INFO.get(reporte_id, {
        'nombre': getattr(CATALOGO.get(reporte_id), 'NOMBRE', reporte_id),
        'template': f'reporte_{reporte_id}.html'
    })
    conn = get_db_connection()
    try:
        usuario_id = session.get('usuario_id')
        rol = session.get('rol', '')
        agentes = _obtener_agentes_activos(conn, usuario_id, rol)
        hoy = datetime.date.today().isoformat()
        
        agente_param = request.args.get('agente', '').strip()
        if not agente_param and agentes:
            online_agents = [a for a in agentes if a['online']]
            agente_param = online_agents[0]['id'] if online_agents else agentes[0]['id']
            
        return render_template(info['template'],
                               reporte_id=reporte_id,
                               reporte_nombre=info['nombre'],
                               agentes=agentes,
                               agente_actual=agente_param,
                               desde=hoy,
                               hasta=hoy,
                               sar_nombre=session.get('nombre', ''))
    finally:
        conn.close()

@bp.route('/admin/api/reporte/<reporte_id>/iniciar', methods=['POST'])
@admin_required
def api_reporte_iniciar(reporte_id):
    if reporte_id not in REPORTES_INFO and reporte_id not in CATALOGO:
        return jsonify({'ok': False, 'error': f"Reporte '{reporte_id}' no existe"}), 404
        
    body = request.get_json(force=True) or {}
    agente = (body.get('agente') or body.get('cliente_id') or '').strip()
    filtros = {k: v for k, v in body.items() if k not in ('agente', 'cliente_id', 'fuente')}
    
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        if not agente:
            return jsonify({'ok': False, 'error': 'No se especificó un agente'}), 400
            
        row = conn.execute("""
            INSERT INTO admin_agent_consultas (sesion_id, tipo, parametros, estado)
            SELECT id, 'ejecutar_reporte', %s::jsonb, 'pendiente'
            FROM admin_agent_sesiones
            WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon')
               OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '')
            ORDER BY ultimo_ping DESC, id DESC
            LIMIT 1
            RETURNING id
        """, (json.dumps({'reporte_id': reporte_id, 'filtros': filtros}), agente, agente, agente)).fetchone()
        
        if not row:
            row = conn.execute("""
                INSERT INTO admin_agent_consultas (sesion_id, tipo, parametros, estado)
                SELECT id, 'ejecutar_reporte', %s::jsonb, 'pendiente'
                FROM admin_agent_sesiones
                WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon')
                   OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '')
                ORDER BY id DESC
                LIMIT 1
                RETURNING id
            """, (json.dumps({'reporte_id': reporte_id, 'filtros': filtros}), agente, agente, agente)).fetchone()
            
        if not row:
            return jsonify({'ok': False, 'error': f"El agente '{agente}' no está registrado"}), 400
            
        conn.commit()
        return jsonify({'ok': True, 'consulta_id': row['id'], 'agente': agente})
    finally:
        conn.close()

@bp.route('/admin/api/reporte/resultado/<int:consulta_id>', methods=['GET'])
@admin_required
def api_reporte_resultado(consulta_id):
    conn = get_db_connection()
    try:
        row = conn.execute("""
            SELECT id, tipo, parametros, respuesta, estado, created_at, respondida_at
            FROM admin_agent_consultas WHERE id = %s
        """, (consulta_id,)).fetchone()
        
        if not row:
            return jsonify({'ok': False, 'error': 'Consulta no encontrada'}), 404
            
        estado = row['estado']
        if estado in ('pendiente', 'procesando'):
            if row['created_at']:
                from datetime import datetime, timezone
                now = datetime.now(timezone.utc)
                created = row['created_at']
                if created.tzinfo is None:
                    created = created.replace(tzinfo=timezone.utc)
                if (now - created).total_seconds() > 120:
                    try:
                        conn.execute("UPDATE admin_agent_consultas SET estado='error', respuesta=%s::jsonb, respondida_at=NOW() WHERE id=%s",
                                     (json.dumps({'error': 'Tiempo de espera agotado (120s). La consulta en el agente no respondió a tiempo.'}), consulta_id))
                        conn.commit()
                    except Exception:
                        pass
                    return jsonify({'ok': False, 'estado': 'error', 'error': 'Tiempo de espera agotado (120s). La consulta tardó más de 2 minutos.'})
            return jsonify({'ok': True, 'estado': estado})
            
        if estado == 'error':
            err = row.get('respuesta')
            err_msg = err.get('error') if isinstance(err, dict) else str(err or 'Error en el agente')
            return jsonify({'ok': False, 'estado': 'error', 'error': err_msg})
            
        if estado == 'lista':
            ans = row['respuesta']
            if isinstance(ans, dict):
                filas = ans.get('filas') if 'filas' in ans else ans.get('rows', [])
            elif isinstance(ans, list):
                filas = ans
            else:
                filas = []
            
            elapsed = 0
            if row['respondida_at'] and row['created_at']:
                elapsed = round((row['respondida_at'] - row['created_at']).total_seconds(), 2)
                
            return jsonify({
                'ok': True,
                'estado': 'lista',
                'filas': filas,
                'rows': filas,
                'total': len(filas),
                'elapsed': elapsed
            })
            
        return jsonify({'ok': True, 'estado': estado})
    finally:
        conn.close()

@bp.route('/admin/api/reporte/<reporte_id>', methods=['POST'])
@bp.route('/api/reporte/<reporte_id>', methods=['POST'])
@admin_required
def ejecutar_reporte_api(reporte_id):
    if reporte_id not in REPORTES_INFO and reporte_id not in CATALOGO:
        return jsonify({'ok': False, 'error': f"Reporte '{reporte_id}' no existe"}), 404
        
    body = request.get_json(force=True) or {}
    agente = (body.get('agente') or body.get('cliente_id') or '').strip()
    
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        if not agente:
            return jsonify({'ok': False, 'error': 'No se especificó un agente remoto'}), 400
            
        filtros = {k: v for k, v in body.items() if k not in ('agente', 'cliente_id', 'fuente')}
        t0 = time.time()
        mod = CATALOGO.get(reporte_id)
        filas = None
        
        # 1. Intentar ejecución remota compilada en el agente solo para reportes estándar
        if reporte_id != 'auditoria_facturas':
            try:
                resp = _ejecutar_consulta_remota(conn, agente, 'ejecutar_reporte', {
                    'reporte_id': reporte_id,
                    'filtros': filtros
                }, timeout=5)
                if isinstance(resp, dict):
                    if 'error' not in resp:
                        if 'filas' in resp:
                            filas = resp['filas']
                        elif 'rows' in resp:
                            filas = resp['rows']
                        else:
                            filas = [resp]
                elif isinstance(resp, list):
                    filas = resp
            except Exception:
                filas = None

        # 2. Si el agente no tiene el reporte compilado o es dinámico, ejecutar vía multi_tabla en el servidor
        if filas is None or not isinstance(filas, list):
            if mod and hasattr(mod, 'tablas_requeridas') and hasattr(mod, 'calcular'):
                tablas = mod.tablas_requeridas(filtros)
                datos = _ejecutar_consulta_remota(conn, agente, 'multi_tabla', {
                    'tablas': tablas
                }, timeout=90)
                if isinstance(datos, dict) and 'error' in datos:
                    return jsonify({'ok': False, 'error': datos['error'], 'filas': [], 'rows': []}), 400

                # Resolver terceros en lote desde la cache central
                if isinstance(datos, dict):
                    if 'REG_CTAS' in datos and isinstance(datos['REG_CTAS'], list):
                        cod_ters = {str(r.get('TERCERO', '') or '').strip() for r in datos['REG_CTAS']}
                        datos['TERCEROS_MAP'] = _resolver_terceros_lote(conn, agente, cod_ters)
                    elif 'PROD_FACT1' in datos and isinstance(datos['PROD_FACT1'], list):
                        cod_ters = {str(r.get('CLIENTE', '') or '').strip() for r in datos['PROD_FACT1']}
                        datos['TERCEROS_MAP'] = _resolver_terceros_lote(conn, agente, cod_ters)

                filas = mod.calcular(datos, filtros)
            else:
                return jsonify({'ok': False, 'error': f"Reporte '{reporte_id}' no pudo ser ejecutado", 'filas': [], 'rows': []}), 500
                
        elapsed = round(time.time() - t0, 2)
        total_cnt = len(filas) if hasattr(filas, '__len__') else 0
        return jsonify({
            'ok': True,
            'rows': filas,
            'filas': filas,
            'total': total_cnt,
            'elapsed': elapsed,
            'agente': agente
        })
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'filas': [], 'rows': []}), 500
    finally:
        conn.close()

@bp.route('/admin/api/reporte/<reporte_id>/excel', methods=['POST'])
@bp.route('/api/reporte/<reporte_id>/excel', methods=['POST'])
@admin_required
def exportar_excel_api(reporte_id):
    try:
        body = request.get_json(force=True) or {}
        agente = (body.get('agente') or body.get('cliente_id') or '').strip()
        filas = body.get('filas') or body.get('rows')
        filtros = {k: v for k, v in body.items() if k not in ('agente', 'cliente_id', 'fuente', 'filas', 'rows')}
        mod = CATALOGO.get(reporte_id)
        
        conn = get_db_connection()
        try:
            if not filas or not isinstance(filas, list):
                if not agente:
                    agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
                    if agentes:
                        agente = agentes[0]['id']
                if not agente:
                    return 'Agente requerido', 400
                if reporte_id != 'auditoria_facturas':
                    try:
                        resp = _ejecutar_consulta_remota(conn, agente, 'ejecutar_reporte', {
                            'reporte_id': reporte_id,
                            'filtros': filtros
                        }, timeout=5)
                        if isinstance(resp, dict) and 'error' not in resp:
                            filas = resp.get('filas') or resp.get('rows') or [resp]
                        elif isinstance(resp, list):
                            filas = resp
                        else:
                            filas = None
                    except Exception:
                        filas = None

                if filas is None or not isinstance(filas, list):
                    if mod and hasattr(mod, 'tablas_requeridas') and hasattr(mod, 'calcular'):
                        tablas = mod.tablas_requeridas(filtros)
                        datos = _ejecutar_consulta_remota(conn, agente, 'multi_tabla', {'tablas': tablas}, timeout=90)
                        if isinstance(datos, dict):
                            if 'REG_CTAS' in datos and isinstance(datos['REG_CTAS'], list):
                                cod_ters = {str(r.get('TERCERO', '') or '').strip() for r in datos['REG_CTAS']}
                                datos['TERCEROS_MAP'] = _resolver_terceros_lote(conn, agente, cod_ters)
                            elif 'PROD_FACT1' in datos and isinstance(datos['PROD_FACT1'], list):
                                cod_ters = {str(r.get('CLIENTE', '') or '').strip() for r in datos['PROD_FACT1']}
                                datos['TERCEROS_MAP'] = _resolver_terceros_lote(conn, agente, cod_ters)
                        filas = mod.calcular(datos, filtros)
                    else:
                        filas = []
        finally:
            conn.close()
            
        if mod and hasattr(mod, 'generar_excel'):
            buf = mod.generar_excel(filas, filtros, 'remoto', agente)
        else:
            buf = _generar_excel_buffer(reporte_id, filas, filtros, mod, fuente='remoto', cliente_id=agente)
            
        filename = f"{reporte_id}_{datetime.date.today().isoformat()}.xlsx"
        return send_file(
            buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


@bp.route('/admin/api/cuentas', methods=['GET', 'POST'])
@bp.route('/api/cuentas', methods=['GET', 'POST'])
@admin_required
def api_cuentas():
    body = request.get_json(silent=True) or {}
    agente = (request.args.get('agente') or body.get('agente') or request.args.get('cliente_id') or body.get('cliente_id') or '').strip()
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        if not agente:
            return jsonify({'ok': False, 'error': 'Agente requerido', 'cuentas': []}), 400
        
        ses_row = conn.execute(
            "SELECT db_version, cuentas_cache FROM admin_agent_sesiones "
            "WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon') OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '') "
            "ORDER BY ultimo_ping DESC, id DESC LIMIT 1",
            (agente, agente, agente)
        ).fetchone()
        db_ver = (ses_row['db_version'] if ses_row and ses_row['db_version'] else '')
        cached_cta = (ses_row['cuentas_cache'] if ses_row and ses_row['cuentas_cache'] else None)

        if cached_cta and isinstance(cached_cta, list) and len(cached_cta) > 0:
            return jsonify({'ok': True, 'cuentas': cached_cta, 'db_version': db_ver, 'total': len(cached_cta), 'origen': 'agente_cache'})

        datos = _ejecutar_consulta_remota(conn, agente, 'multi_tabla', {
            'tablas': [{'tabla': 'CUENTAS', 'campos': ['CODIGO', 'NOMBRE', 'TIPO'], 'filtros': {}}]
        }, timeout=30)
        
        raw_cuentas = datos.get('CUENTAS', []) if isinstance(datos, dict) else []
        cuentas = []
        for r in raw_cuentas:
            c = str(r.get('CODIGO', '') or '').strip()
            n = str(r.get('NOMBRE', '') or '').strip()
            t = str(r.get('TIPO', '') or '').strip().upper()
            if c:
                cuentas.append({'codigo': c, 'nombre': n, 'tipo': t, 'es_movimiento': (t == 'D')})
        cuentas.sort(key=lambda x: x['codigo'])

        if cuentas:
            try:
                conn.execute(
                    "UPDATE admin_agent_sesiones SET cuentas_cache=%s::jsonb "
                    "WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon') OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '')",
                    (json.dumps(cuentas), agente, agente, agente)
                )
                conn.commit()
            except Exception:
                pass

        return jsonify({'ok': True, 'cuentas': cuentas, 'db_version': db_ver, 'total': len(cuentas)})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'cuentas': []}), 500
    finally:
        conn.close()


@bp.route('/admin/api/empresas', methods=['GET', 'POST'])
@bp.route('/api/empresas', methods=['GET', 'POST'])
@admin_required
def api_empresas():
    body = request.get_json(silent=True) or {}
    agente = (request.args.get('agente') or body.get('agente') or request.args.get('cliente_id') or body.get('cliente_id') or '').strip()
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        if not agente:
            return jsonify({'ok': False, 'error': 'Agente requerido', 'empresas': []}), 400
        
        ses_row = conn.execute(
            "SELECT db_version, empresas_cache FROM admin_agent_sesiones "
            "WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon') OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '') "
            "ORDER BY ultimo_ping DESC, id DESC LIMIT 1",
            (agente, agente, agente)
        ).fetchone()
        db_ver = (ses_row['db_version'] if ses_row and ses_row['db_version'] else '')
        cached_emp = (ses_row['empresas_cache'] if ses_row and ses_row['empresas_cache'] else None)
        
        if cached_emp and isinstance(cached_emp, list) and len(cached_emp) > 0:
            return jsonify({'ok': True, 'empresas': cached_emp, 'db_version': db_ver, 'total': len(cached_emp), 'origen': 'agente_cache'})

        datos = _ejecutar_consulta_remota(conn, agente, 'multi_tabla', {
            'tablas': [{'tabla': 'EMPRESAS', 'campos': ['COD_EMP', 'NOM_EMP', 'NIT'], 'filtros': {}}]
        }, timeout=30)
        
        raw_emp = datos.get('EMPRESAS', []) if isinstance(datos, dict) else []
        empresas = []
        for r in raw_emp:
            cod = str(r.get('COD_EMP', '') or '').strip()
            nom = str(r.get('NOM_EMP', '') or '').strip()
            nit = str(r.get('NIT', '') or '').strip()
            if cod or nom:
                empresas.append({'codigo': cod, 'nombre': nom, 'nit': nit})
        
        empresas.sort(key=lambda x: x['codigo'] if x['codigo'] else x['nombre'])
        
        if empresas:
            try:
                conn.execute(
                    "UPDATE admin_agent_sesiones SET empresas_cache=%s::jsonb "
                    "WHERE LOWER(cliente_id) IN (LOWER(%s), LOWER(%s) || '_daemon') OR LOWER(cliente_id) = REPLACE(LOWER(%s), '_daemon', '')",
                    (json.dumps(empresas), agente, agente, agente)
                )
                conn.commit()
            except Exception:
                pass
                
        return jsonify({'ok': True, 'empresas': empresas, 'db_version': db_ver, 'total': len(empresas)})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'empresas': []}), 500
    finally:
        conn.close()


@bp.route('/admin/api/tipos_doc', methods=['GET', 'POST'])
@bp.route('/api/tipos_doc', methods=['GET', 'POST'])
@admin_required
def api_tipos_doc():
    body = request.get_json(silent=True) or {}
    agente = (request.args.get('agente') or body.get('agente') or request.args.get('cliente_id') or body.get('cliente_id') or '').strip()
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        if not agente:
            return jsonify({'ok': False, 'error': 'Agente requerido', 'tipos': [], 'por_empresa': {}}), 400
        
        datos = _ejecutar_consulta_remota(conn, agente, 'multi_tabla', {
            'tablas': [
                {'tabla': 'TIPO_DOC', 'campos': ['CODIGO', 'NOMBRE', 'TIPO_INVE', 'MANEJAINVE'], 'filtros': {}},
                {'tabla': 'allegra_config', 'campos': ['EMPRESA', 'TIP_DOC', 'NUM_INICIO'], 'filtros': {}},
                {'tabla': 'alegra_tiposdoc', 'campos': ['EMPRESA', 'TIP_ADMIN', 'TIP_ALEGRA'], 'filtros': {}},
                {'tabla': 'contabilidad_documentos_contables_configurar', 'campos': ['EMPRESA', 'DOCUMENTO'], 'filtros': {}},
            ]
        }, timeout=30)
        
        raw_tipos   = datos.get('TIPO_DOC', []) if isinstance(datos, dict) else []
        raw_cfg_al  = datos.get('allegra_config', []) if isinstance(datos, dict) and isinstance(datos.get('allegra_config'), list) else []
        raw_td_al   = datos.get('alegra_tiposdoc', []) if isinstance(datos, dict) and isinstance(datos.get('alegra_tiposdoc'), list) else []
        raw_cfg_ct  = datos.get('contabilidad_documentos_contables_configurar', []) if isinstance(datos, dict) and isinstance(datos.get('contabilidad_documentos_contables_configurar'), list) else []
        
        tipos = []
        tipos_dict = {}
        for r in raw_tipos:
            cod = str(r.get('CODIGO', '') or '').strip().upper()
            nom = str(r.get('NOMBRE', '') or '').strip()
            tip_inv = str(r.get('TIPO_INVE', '') or '').strip()
            man_inv = str(r.get('MANEJAINVE', '') or '').strip()
            if cod or nom:
                t_obj = {'codigo': cod, 'nombre': nom, 'tipo_inve': tip_inv, 'maneja_inve': man_inv}
                tipos.append(t_obj)
                tipos_dict[cod] = t_obj
        
        tipos.sort(key=lambda x: x['codigo'])

        # Mapeo dinámico por empresa
        por_empresa = {}
        for r in raw_cfg_al:
            emp = str(r.get('EMPRESA', '') or '').strip().upper()
            tip = str(r.get('TIP_DOC', '') or '').strip().upper()
            num_ini = str(r.get('NUM_INICIO', '') or '').strip()
            if emp:
                if emp not in por_empresa:
                    por_empresa[emp] = {'default_tipo': '', 'num_inicio': '', 'tipos_interfaz': set(), 'tipos_contables': set()}
                if tip:
                    por_empresa[emp]['default_tipo'] = tip
                    por_empresa[emp]['tipos_interfaz'].add(tip)
                if num_ini:
                    por_empresa[emp]['num_inicio'] = num_ini

        for r in raw_td_al:
            emp = str(r.get('EMPRESA', '') or '').strip().upper()
            tip = str(r.get('TIP_ADMIN', '') or '').strip().upper()
            if emp and tip:
                if emp not in por_empresa:
                    por_empresa[emp] = {'default_tipo': '', 'num_inicio': '', 'tipos_interfaz': set(), 'tipos_contables': set()}
                por_empresa[emp]['tipos_interfaz'].add(tip)

        for r in raw_cfg_ct:
            emp = str(r.get('EMPRESA', '') or '').strip().upper()
            doc = str(r.get('DOCUMENTO', '') or '').strip().upper()
            if emp and doc:
                if emp not in por_empresa:
                    por_empresa[emp] = {'default_tipo': '', 'num_inicio': '', 'tipos_interfaz': set(), 'tipos_contables': set()}
                por_empresa[emp]['tipos_contables'].add(doc)

        for emp, d in por_empresa.items():
            d['tipos_interfaz'] = sorted(list(d['tipos_interfaz']))
            d['tipos_contables'] = sorted(list(d['tipos_contables']))

        return jsonify({'ok': True, 'tipos': tipos, 'por_empresa': por_empresa, 'total': len(tipos)})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'tipos': [], 'por_empresa': {}}), 500
    finally:
        conn.close()


@bp.route('/admin/api/reporte/auditoria_facturas/alegra_check', methods=['POST'])
@bp.route('/api/reporte/auditoria_facturas/alegra_check', methods=['POST'])
@admin_required
def api_alegra_check():
    """Consulta facturas puntuales en la API de Alegra para diagnosticar inconsistencias o faltantes."""
    import requests, base64
    body = request.get_json(force=True) or {}
    numeros = body.get('numeros') or []
    empresa = str(body.get('empresa', '') or '').strip().upper()

    ALEGRA_CREDS = {
        'LP': {
            'email': 'electronicajyp@hotmail.com',
            'token': 'aabde447e95a29efb773',
            'nombre': 'ELECTRONICAS J&P',
        },
        '02': {
            'email': 'electronicastvyvideo@hotmail.com',
            'token': 'ade8e319ce85985fb47c',
            'nombre': 'ELECTRONICAS TV & VIDEO',
        },
    }

    creds_list = [ALEGRA_CREDS[empresa]] if empresa in ALEGRA_CREDS else list(ALEGRA_CREDS.values())

    resultados = {}
    for num in numeros[:50]:
        num_str = str(num).strip()
        if not num_str:
            continue
        encontrado = None
        for cred in creds_list:
            auth_str = base64.b64encode(f"{cred['email']}:{cred['token']}".encode()).decode()
            headers = {
                'Authorization': f'Basic {auth_str}',
                'Accept': 'application/json',
            }
            try:
                url = f"https://api.alegra.com/api/v1/invoices?number={num_str}&limit=5"
                r = requests.get(url, headers=headers, timeout=8)
                if r.ok:
                    items = r.json()
                    for inv in items:
                        num_doc = str(inv.get('numberTemplate', {}).get('fullNumber') or inv.get('id') or inv.get('number') or '')
                        if num_str in num_doc or str(inv.get('number', '')) == num_str:
                            encontrado = {
                                'id': inv.get('id'),
                                'numero': num_doc or inv.get('number'),
                                'fecha': inv.get('date'),
                                'cliente': (inv.get('client', {}) or {}).get('name'),
                                'cliente_nit': (inv.get('client', {}) or {}).get('identification'),
                                'total': float(inv.get('total', 0) or 0),
                                'subtotal': float(inv.get('subtotal', 0) or 0),
                                'total_pagado': float(inv.get('totalPaid', 0) or 0),
                                'estado': inv.get('status'),
                                'empresa_alegra': cred['nombre'],
                                'items_count': len(inv.get('items', [])),
                            }
                            break
                if encontrado:
                    break
            except Exception:
                pass

        if encontrado:
            resultados[num_str] = {'ok': True, 'existe': True, 'alegra': encontrado}
        else:
            resultados[num_str] = {'ok': True, 'existe': False, 'mensaje': 'No encontrada en Alegra'}

    return jsonify({'ok': True, 'resultados': resultados})


@bp.route('/admin/api/terceros/stats', methods=['GET', 'POST'])
@bp.route('/api/terceros/stats', methods=['GET', 'POST'])
@admin_required
def api_terceros_stats():
    body = request.get_json(silent=True) or {}
    agente = (request.args.get('agente') or body.get('agente') or request.args.get('cliente_id') or body.get('cliente_id') or '').strip()
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        agente_clean = agente.strip().lower().replace('_daemon', '')
        
        conn.execute("""
            CREATE TABLE IF NOT EXISTS admin_terceros_cache (
                id SERIAL PRIMARY KEY,
                cliente_id VARCHAR(100) NOT NULL,
                cod_ter VARCHAR(50) NOT NULL,
                nombre VARCHAR(250),
                nit VARCHAR(50),
                updated_at TIMESTAMPTZ DEFAULT NOW(),
                UNIQUE(cliente_id, cod_ter)
            )
        """)
        conn.commit()

        row = conn.execute("SELECT COUNT(*) AS total FROM admin_terceros_cache WHERE LOWER(cliente_id) = LOWER(%s)", (agente_clean,)).fetchone()
        total = row['total'] if row else 0
        return jsonify({'ok': True, 'total': total, 'cliente_id': agente_clean})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'total': 0}), 500
    finally:
        conn.close()


@bp.route('/admin/api/terceros/sync', methods=['POST'])
@bp.route('/api/terceros/sync', methods=['POST'])
@admin_required
def api_terceros_sync():
    body = request.get_json(silent=True) or {}
    agente = (request.args.get('agente') or body.get('agente') or request.args.get('cliente_id') or body.get('cliente_id') or '').strip()
    conn = get_db_connection()
    try:
        if not agente:
            agentes = _obtener_agentes_activos(conn, session.get('usuario_id'), session.get('rol', ''))
            if agentes:
                agente = agentes[0]['id']
        if not agente:
            return jsonify({'ok': False, 'error': 'Agente requerido'}), 400
        
        # Forzar sincronización completa de la tabla TERCEROS
        res = _resolver_terceros_lote(conn, agente, ['_FORCE_ALL_SYNC_'])
        agente_clean = agente.strip().lower().replace('_daemon', '')
        row = conn.execute("SELECT COUNT(*) AS total FROM admin_terceros_cache WHERE LOWER(cliente_id) = LOWER(%s)", (agente_clean,)).fetchone()
        total = row['total'] if row else 0
        return jsonify({'ok': True, 'total': total, 'cliente_id': agente_clean})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500
    finally:
        conn.close()

