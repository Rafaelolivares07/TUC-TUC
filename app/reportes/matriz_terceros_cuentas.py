"""reportes/matriz_terceros_cuentas.py -- Matriz de netos por tercero y cuenta."""

NOMBRE    = 'Matriz Terceros por Cuentas'
ID        = 'matriz_terceros_cuentas'
CATEGORIA = 'Contabilidad'

FILTROS_UI = [
    {'id': 'empresa', 'label': 'Empresa',   'tipo': 'select',      'requerido': False},
    {'id': 'desde',   'label': 'Desde',     'tipo': 'date',        'requerido': True},
    {'id': 'hasta',   'label': 'Hasta',     'tipo': 'date',        'requerido': True},
    {'id': 'tipos',   'label': 'Tipos doc', 'tipo': 'multiselect', 'requerido': False},
]


def tablas_requeridas(filtros):
    if filtros.get('tipos_init'):
        return [{'tabla': 'TIPO_DOC', 'campos': ['CODIGO', 'NOMBRE'], 'filtros': {}}]

    empresa = str(filtros.get('empresa', '') or '').strip().upper()
    desde   = filtros.get('desde', '')
    hasta   = filtros.get('hasta', '')

    def a_lapso(f):
        return f.replace('-', '') if f else ''

    filtros_rc = {}
    if empresa:
        filtros_rc['EMPRESA'] = empresa
    if desde or hasta:
        filtros_rc['LAPSO'] = {}
        if desde:
            filtros_rc['LAPSO']['desde'] = a_lapso(desde)
        if hasta:
            filtros_rc['LAPSO']['hasta'] = a_lapso(hasta)

    return [
        {
            'tabla':   'REG_CTAS',
            'campos':  ['TERCERO', 'CUENTA', 'TOT_DEB', 'TOT_CRE', 'EMPRESA', 'LAPSO', 'TIPO'],
            'filtros': filtros_rc,
        },
        {
            'tabla':   'TERCEROS',
            'campos':  ['COD_TER', 'NOMBRE', 'IDENTIFICA', 'NIT'],
            'filtros': {},
        },
        {
            'tabla':   'CUENTAS',
            'campos':  ['CODIGO', 'NOMBRE'],
            'filtros': {},
        },
    ]


def calcular(datos, filtros):
    if filtros.get('tipos_init'):
        tipos = []
        for r in datos.get('TIPO_DOC', []):
            cod = str(r.get('CODIGO', '') or '').strip()
            nom = str(r.get('NOMBRE', '') or '').strip()
            if cod:
                tipos.append({'codigo': cod, 'nombre': nom or cod})
        tipos.sort(key=lambda x: x['codigo'])
        return tipos

    tipos_sel = filtros.get('tipos', [])
    if isinstance(tipos_sel, str):
        tipos_sel = [tipos_sel] if tipos_sel else []
    tipos_sel = {t.upper().strip() for t in tipos_sel if t}

    terceros = {}
    for r in datos.get('TERCEROS', []):
        cod = str(r.get('COD_TER', '') or '').strip()
        if not cod:
            continue
        terceros[cod] = {
            'identificacion': str(r.get('IDENTIFICA', '') or r.get('NIT', '') or '').strip(),
            'nombre':         str(r.get('NOMBRE', '') or '').strip() or cod,
        }

    cuentas_nom = {
        str(r.get('CODIGO', '') or '').strip(): str(r.get('NOMBRE', '') or '').strip()
        for r in datos.get('CUENTAS', [])
        if str(r.get('CODIGO', '') or '').strip()
    }

    matriz = {}
    cuentas_con_mov = set()

    for r in datos.get('REG_CTAS', []):
        if tipos_sel:
            tipo = str(r.get('TIPO', '') or '').strip().upper()
            if tipo not in tipos_sel:
                continue

        cod_ter = str(r.get('TERCERO', '') or '').strip()
        cuenta  = str(r.get('CUENTA', '') or '').strip()
        if not cod_ter or not cuenta:
            continue

        neto = float(r.get('TOT_DEB', 0) or 0) - float(r.get('TOT_CRE', 0) or 0)
        if abs(neto) < 0.005:
            continue

        if cod_ter not in matriz:
            ter = terceros.get(cod_ter, {})
            matriz[cod_ter] = {
                'identificacion': ter.get('identificacion', ''),
                'nombre':         ter.get('nombre', cod_ter),
                'valores':        {},
            }
        valores = matriz[cod_ter]['valores']
        valores[cuenta] = valores.get(cuenta, 0.0) + neto

    for fila in matriz.values():
        fila['valores'] = {
            cta: round(val, 2)
            for cta, val in fila['valores'].items()
            if abs(val) >= 0.005
        }
        cuentas_con_mov.update(fila['valores'].keys())

    cuentas = [
        {'codigo': cta, 'nombre': cuentas_nom.get(cta, '')}
        for cta in sorted(cuentas_con_mov)
    ]

    rows = []
    for fila in matriz.values():
        if not fila['valores']:
            continue
        rows.append(fila)

    rows.sort(key=lambda r: (r.get('nombre', '').upper(), r.get('identificacion', '')))
    if rows:
        rows[0]['__cuentas'] = cuentas
    return rows


def generar_excel(rows, filtros, fuente, cliente_id):
    import io
    import datetime
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Matriz Terceros'[:31]

    azul = '003F7F'
    gris = 'F0F4FF'
    tit_font = Font(bold=True, size=12, color='FFFFFF')
    hdr_font = Font(bold=True, color='FFFFFF')
    meta_font = Font(italic=True, color='555555', size=10)
    fill_azul = PatternFill('solid', fgColor=azul)
    fill_gris = PatternFill('solid', fgColor=gris)

    fila = 1
    c = ws.cell(fila, 1, NOMBRE)
    c.font = tit_font
    c.fill = fill_azul
    fila += 1

    fuente_txt = f'Fuente: {(fuente or "local").upper()}'
    if cliente_id:
        fuente_txt += f'  -  {cliente_id}'
    ws.cell(fila, 1, fuente_txt).font = meta_font
    fila += 1

    filtros_lbl = {f['id']: f['label'] for f in FILTROS_UI}
    for key in ('empresa', 'desde', 'hasta', 'tipos'):
        val = filtros.get(key, '')
        if isinstance(val, list):
            val = ', '.join(str(v) for v in val if v)
        if val:
            ws.cell(fila, 1, f'{filtros_lbl.get(key, key)}: {val}').font = meta_font
            fila += 1

    fila += 1

    cuentas = rows[0].get('__cuentas', []) if rows else []
    headers = ['Identificacion', 'Nombre'] + [
        f"{c['codigo']} - {c['nombre']}" if c.get('nombre') else c['codigo']
        for c in cuentas
    ]
    for ci, label in enumerate(headers, 1):
        cell = ws.cell(fila, ci, label)
        cell.font = hdr_font
        cell.fill = fill_azul
    fila += 1

    if not rows:
        ws.cell(fila, 1, 'Sin resultados').font = meta_font
    else:
        for ri, r in enumerate(rows):
            fill = fill_gris if ri % 2 == 1 else None
            ws.cell(fila, 1, r.get('identificacion', ''))
            ws.cell(fila, 2, r.get('nombre', ''))
            for ci, cta in enumerate(cuentas, 3):
                val = r.get('valores', {}).get(cta['codigo'], 0)
                if val:
                    cell = ws.cell(fila, ci, val)
                    cell.number_format = '#,##0.00'
            if fill:
                for ci in range(1, len(headers) + 1):
                    ws.cell(fila, ci).fill = fill
            fila += 1

    fila += 1
    ws.cell(fila, 1, f'Generado: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}').font = meta_font

    ws.freeze_panes = 'C7'
    for col in ws.columns:
        letra = col[0].column_letter
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        ws.column_dimensions[letra].width = min(max_len + 2, 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
