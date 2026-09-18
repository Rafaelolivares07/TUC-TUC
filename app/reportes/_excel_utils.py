"""reportes/_excel_utils.py -- Helper compartido para generar Excel con encabezado."""
import io
import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def generar_excel(nombre_reporte, filtros_ui, excel_cols, rows, filtros, fuente, cliente_id):
    """
    Genera un Excel con:
    - Fila de título (nombre del reporte)
    - Fila de fuente (local/remoto + cliente)
    - Filas de filtros aplicados (usando filtros_ui como etiquetas)
    - Encabezados de columnas (usando excel_cols: [(key, label), ...])
    - Filas de datos con formato numérico automático
    - Pie con timestamp
    """
    if not HAS_OPENPYXL:
        raise RuntimeError('openpyxl no instalado')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_reporte[:31]

    AZ  = '003F7F'
    GR  = 'F0F4FF'
    tit_font  = Font(bold=True, size=12, color='FFFFFF')
    tit_fill  = PatternFill('solid', fgColor=AZ)
    hdr_font  = Font(bold=True, color='FFFFFF')
    hdr_fill  = PatternFill('solid', fgColor=AZ)
    meta_font = Font(italic=True, color='555555', size=10)
    pie_font  = Font(italic=True, color='888888', size=9)
    num_fmt   = '#,##0.00'

    fila = 1

    # Título
    c = ws.cell(fila, 1, nombre_reporte)
    c.font = tit_font
    c.fill = tit_fill
    fila += 1

    # Fuente
    fuente_txt = f'Fuente: {(fuente or "local").upper()}'
    if cliente_id:
        fuente_txt += f'  —  {cliente_id}'
    ws.cell(fila, 1, fuente_txt).font = meta_font
    fila += 1

    # Filtros aplicados
    for f_ui in (filtros_ui or []):
        fid = f_ui['id']
        val = filtros.get(fid, '')
        if isinstance(val, list):
            val = ', '.join(str(v) for v in val if v)
        if val:
            ws.cell(fila, 1, f"{f_ui['label']}: {val}").font = meta_font
            fila += 1

    fila += 1  # línea en blanco

    if not rows:
        ws.cell(fila, 1, 'Sin resultados').font = meta_font
    else:
        keys   = [k  for k, _ in excel_cols]
        labels = [lb for _, lb in excel_cols]

        # Encabezados
        for ci, label in enumerate(labels, 1):
            c = ws.cell(fila, ci, label)
            c.font = hdr_font
            c.fill = hdr_fill
        fila += 1

        # Detectar columnas numéricas
        sample = rows[0]
        num_keys = {k for k in keys if isinstance(sample.get(k), (int, float))}

        # Datos
        for ri, r in enumerate(rows):
            fill_par = PatternFill('solid', fgColor=GR) if ri % 2 == 1 else None
            for ci, k in enumerate(keys, 1):
                val = r.get(k, '')
                c = ws.cell(fila, ci, val)
                if k in num_keys:
                    c.number_format = num_fmt
                if fill_par:
                    c.fill = fill_par
            fila += 1

    # Pie
    fila += 1
    ws.cell(fila, 1, f'Generado: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}').font = pie_font

    # Ancho automático de columnas
    for col in ws.columns:
        letra = col[0].column_letter
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        ws.column_dimensions[letra].width = min(max_len + 2, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
